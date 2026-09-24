"""Compare exported Excel formula results with the system's cached values."""

import re
import sys

import openpyxl
from openpyxl.utils.cell import range_boundaries


if len(sys.argv) != 2:
    raise SystemExit("Usage: verify_export_formulas.py <export.xlsx>")

formulas = openpyxl.load_workbook(sys.argv[1], data_only=False)
values = openpyxl.load_workbook(sys.argv[1], data_only=True)


def number(sheet, reference):
    value = values[sheet][reference].value
    return 0.0 if value is None else float(value)


def evaluate(sheet, text):
    expression = text[1:].replace("$", "")
    guard = re.fullmatch(r"IF\((K\d+)=0,0,(.*)\)", expression)
    if guard:
        if number(sheet, guard.group(1)) == 0:
            return 0.0
        expression = guard.group(2)

    def sum_range(address):
        start_col, start_row, end_col, end_row = range_boundaries(address)
        return sum(
            float(cell.value or 0)
            for row in values[sheet].iter_rows(
                min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col
            )
            for cell in row
        )

    ranges = []

    def hold_range(match):
        ranges.append(match.group(1))
        return f"__RANGE{len(ranges) - 1}__"

    expression = re.sub(r"SUM\(([A-Z]+\d+:[A-Z]+\d+)\)", hold_range, expression)
    expression = expression.replace("SUM(", "sum_items(")
    cross = []

    def hold_cross(match):
        cross.append((match.group(1), match.group(2)))
        return f"__CROSS{len(cross) - 1}__"

    expression = re.sub(r"([A-Z]+)!([A-Z]+\d+)", hold_cross, expression)
    expression = re.sub(r"\b([A-Z]{1,2}\d+)\b", lambda m: f'number(sheet,"{m.group(1)}")', expression)
    for index, (other_sheet, reference) in enumerate(cross):
        expression = expression.replace(f"__CROSS{index}__", f'number("{other_sheet}","{reference}")')
    for index, address in enumerate(ranges):
        expression = expression.replace(f"__RANGE{index}__", f'sum_range("{address}")')
    if not re.fullmatch(r'[A-Za-z0-9_(),.:"+\-*/ ]+', expression):
        raise ValueError(f"Unexpected formula: {text}")
    try:
        return eval(expression, {"__builtins__": {}}, {
            "number": number, "sheet": sheet, "sum_range": sum_range,
            "sum_items": lambda *items: sum(items),
        })
    except Exception as error:
        raise ValueError(f"Cannot evaluate {sheet}!{text} as {expression}") from error


mismatches = []
formula_count = 0
for sheet in formulas:
    for row in sheet.iter_rows(min_row=9, max_row=169, min_col=6, max_col=19):
        for cell in row:
            if cell.data_type != "f":
                continue
            formula_count += 1
            expected = number(sheet.title, cell.coordinate)
            actual = evaluate(sheet.title, cell.value)
            tolerance = max(0.02, abs(expected) * 1e-12)
            if abs(actual - expected) > tolerance:
                mismatches.append((sheet.title, cell.coordinate, cell.value, expected, actual))

print("Checked", formula_count, "formulas; mismatches:", len(mismatches))
for mismatch in mismatches[:40]:
    print(mismatch)
if mismatches:
    raise SystemExit(1)
