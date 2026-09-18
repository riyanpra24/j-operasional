"""Read-only comparison of the current RKA workbook budget block and active system data."""
from decimal import Decimal
import json
import re
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook


path = Path(sys.argv[1])
schema = json.loads(Path("app/Config/RkaTemplate2026.json").read_text(encoding="utf-8"))
units = ["Korporat Kanwil", "Kanwil", "Surabaya", "Kediri", "Malang", "Madiun", "Banyuwangi"]
source_units = units[1:]

php = r'''
define("FCPATH",realpath("public").DIRECTORY_SEPARATOR);
define("ENVIRONMENT","development");
require "app/Config/Paths.php";
$paths=new Config\Paths();require $paths->systemDirectory."/Boot.php";
CodeIgniter\Boot::bootConsole($paths);
$service=new App\Libraries\RkaBudgetService();$out=[];
foreach(App\Libraries\RkaCalculator::UNITS as $unit){
    $row=$service->find($unit,2026);
    $out[$unit]=$row===null?null:[
        "inputs"=>json_decode($row["inputs_json"],true),
        "calculated"=>json_decode($row["calculated_json"],true),
        "source_name"=>$row["source_name"],"source_hash"=>$row["source_hash"]
    ];
}
echo json_encode($out,JSON_THROW_ON_ERROR);
'''
system = json.loads(subprocess.check_output(["C:/xampp/php/php.exe", "-r", php], text=True))
formulas = load_workbook(path, data_only=False, read_only=False)
values = load_workbook(path, data_only=True, read_only=False)


def decimal(value):
    if value in (None, "", "-"):
        return Decimal(0)
    return Decimal(str(value))


def normalize(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def normalize_formula(value):
    return re.sub(r"[\s$=]", "", str(value or "")).upper()


def shift_formula(value):
    return re.sub(
        r"(?<![A-Z0-9_])([C-H])(\d+)",
        lambda match: chr(ord(match.group(1)) + 3) + match.group(2),
        value,
    )


report = {"units": {}, "totals": {}}
for unit in units:
    sheet_name = unit.upper()
    ws_formula = formulas[sheet_name]
    ws_value = values[sheet_name]
    stored = system[unit]
    unit_report = {
        "input_mismatches": [],
        "calculated_mismatches": [],
        "label_mismatches": [],
        "header_mismatches": [],
        "formula_mismatches": [],
        "missing_budget_formulas": [],
        "source_input_formula_cells": [],
        "invalid_corporate_input_formulas": [],
    }
    for source_column, workbook_column in zip("CDEFGH", "FGHIJK"):
        expected = schema["columns"][source_column]
        actual = ws_value[f"{workbook_column}6"].value
        if normalize(actual) != normalize(expected):
            unit_report["header_mismatches"].append([f"{workbook_column}6", actual, expected])
    for row_text, definition in schema["rows"].items():
        row = int(row_text)
        label = ws_value[f"E{row}"].value or ws_value[f"D{row}"].value
        if normalize(label) != normalize(definition["label"]):
            unit_report["label_mismatches"].append([row, label, definition["label"]])
    if stored is None:
        unit_report["missing_system_record"] = True
        report["units"][unit] = unit_report
        continue
    for row in schema["input_rows"]:
        for source_column, workbook_column in zip("CDEFG", "FGHIJ"):
            system_cell = f"{source_column}{row}"
            workbook_cell = f"{workbook_column}{row}"
            excel_amount = decimal(ws_value[workbook_cell].value)
            system_amount = decimal(stored["inputs"][system_cell])
            if excel_amount != system_amount:
                unit_report["input_mismatches"].append([workbook_cell, str(excel_amount), system_cell, str(system_amount)])
            formula = ws_formula[workbook_cell].value
            if isinstance(formula, str) and formula.startswith("="):
                if unit == "Korporat Kanwil":
                    actual_terms = normalize_formula(formula).replace("'", "").split("+")
                    expected_terms = [normalize_formula(f"{source.upper()}!{workbook_cell}") for source in source_units]
                    if sorted(actual_terms) != sorted(expected_terms):
                        unit_report["invalid_corporate_input_formulas"].append([workbook_cell, formula])
                else:
                    unit_report["source_input_formula_cells"].append([workbook_cell, formula])
    for row_text in schema["rows"]:
        row = int(row_text)
        for source_column, workbook_column in zip("CDEFGH", "FGHIJK"):
            system_cell = f"{source_column}{row}"
            workbook_cell = f"{workbook_column}{row}"
            excel_amount = decimal(ws_value[workbook_cell].value)
            system_amount = decimal(stored["calculated"][system_cell])
            if excel_amount != system_amount:
                unit_report["calculated_mismatches"].append([workbook_cell, str(excel_amount), system_cell, str(system_amount)])
    for source_cell, source_formula in schema["formulas"].items():
        match = re.fullmatch(r"([C-H])(\d+)", source_cell)
        workbook_cell = f"{chr(ord(match.group(1)) + 3)}{match.group(2)}"
        actual = ws_formula[workbook_cell].value
        expected = shift_formula(source_formula)
        if actual in (None, ""):
            unit_report["missing_budget_formulas"].append(workbook_cell)
        elif normalize_formula(actual) != normalize_formula(expected):
            unit_report["formula_mismatches"].append([workbook_cell, actual, expected])
    unit_report["excel_H166"] = str(decimal(ws_value["K166"].value))
    unit_report["system_H166"] = stored["calculated"]["H166"]
    report["units"][unit] = unit_report

report["totals"] = {
    key: sum(len(unit_report.get(key, [])) for unit_report in report["units"].values())
    for key in (
        "input_mismatches", "calculated_mismatches", "label_mismatches", "header_mismatches",
        "formula_mismatches", "missing_budget_formulas", "source_input_formula_cells",
        "invalid_corporate_input_formulas",
    )
}
print(json.dumps(report, ensure_ascii=False, indent=2))
