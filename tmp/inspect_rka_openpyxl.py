import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def main():
    rka_path = Path(sys.argv[1])
    lr_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])

    rka = load_workbook(rka_path, data_only=False, read_only=False)
    rka_values = load_workbook(rka_path, data_only=True, read_only=False)
    lr = load_workbook(lr_path, data_only=True, read_only=False)

    lr_accounts = {}
    for ws in lr.worksheets:
        for row in ws.iter_rows():
            entity = row[0].value if len(row) > 0 else None
            account = row[2].value if len(row) > 2 else None
            description = clean(row[3].value if len(row) > 3 else None)
            amount = row[7].value if len(row) > 7 else None
            if not entity or account is None or not isinstance(amount, (int, float)):
                continue
            code_match = re.search(r"\d{5,12}", clean(account))
            if not code_match:
                continue
            code = code_match.group(0)
            item = lr_accounts.setdefault(code, {"descriptions": set(), "total": 0.0, "by_sheet": {}})
            if description:
                item["descriptions"].add(description)
            item["total"] += amount
            item["by_sheet"][ws.title] = item["by_sheet"].get(ws.title, 0.0) + amount

    result = {
        "rka_file": str(rka_path),
        "lr_file": str(lr_path),
        "defined_names": [],
        "external_links": [],
        "sheets": [],
        "lr_account_count": len(lr_accounts),
        "account_matches": [],
    }

    for name, defined in rka.defined_names.items():
        result["defined_names"].append({"name": name, "attr_text": defined.attr_text})
    for link in getattr(rka, "_external_links", []):
        result["external_links"].append({"file_link": getattr(getattr(link, "file_link", None), "Target", None)})

    seen_account_matches = set()
    for ws in rka.worksheets:
        ws_values = rka_values[ws.title]
        formulas = []
        formula_functions = Counter()
        external_formulas = []
        cross_sheet_formulas = []
        nonempty_rows = []
        keyword_rows = []
        account_cells = []
        keyword_pattern = re.compile(
            r"anggaran|realisasi|rka|rkap|persen|prosentase|persentase|sisa|varian|variance|ytd|"
            r"laba|rugi|pendapatan|beban|biaya|surplus|defisit|saldo|akun|rekening",
            re.I,
        )
        for row in ws.iter_rows():
            cells = []
            row_text = []
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                cells.append({"cell": cell.coordinate, "value": value})
                row_text.append(clean(value))
                text = clean(value)
                if isinstance(value, str) and value.startswith("="):
                    cached = ws_values[cell.coordinate].value
                    formula = {"cell": cell.coordinate, "formula": value, "cached": cached}
                    formulas.append(formula)
                    funcs = re.findall(r"(?<![A-Z0-9_])([A-Z][A-Z0-9._]*)\s*\(", value.upper())
                    if funcs:
                        formula_functions.update(funcs)
                    else:
                        formula_functions["ARITHMETIC"] += 1
                    if re.search(r"\[[^\]]+\]|\.xlsx|\.xls|\\|https?://", value, re.I):
                        external_formulas.append(formula)
                    if "!" in value:
                        cross_sheet_formulas.append(formula)
                for code in re.findall(r"\b\d{5,12}\b", text):
                    if code in lr_accounts:
                        account_cells.append({"cell": cell.coordinate, "account": code, "value": value})
                        key = (ws.title, cell.coordinate, code)
                        if key not in seen_account_matches:
                            seen_account_matches.add(key)
                            lr_item = lr_accounts[code]
                            result["account_matches"].append({
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "account": code,
                                "rka_value": value,
                                "row_values": [c.value for c in row],
                                "lr_descriptions": sorted(lr_item["descriptions"]),
                                "lr_total": lr_item["total"],
                                "lr_by_sheet": lr_item["by_sheet"],
                            })
            if cells:
                row_record = {"row": row[0].row, "cells": cells}
                if len(nonempty_rows) < 45:
                    nonempty_rows.append(row_record)
                if keyword_pattern.search(" | ".join(row_text)) and len(keyword_rows) < 220:
                    keyword_rows.append(row_record)

        result["sheets"].append({
            "name": ws.title,
            "state": ws.sheet_state,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "merged_ranges": [str(rng) for rng in ws.merged_cells.ranges],
            "formula_count": len(formulas),
            "formula_functions": dict(formula_functions.most_common()),
            "formula_samples": formulas[:160],
            "external_formulas": external_formulas[:200],
            "cross_sheet_formulas": cross_sheet_formulas[:200],
            "first_nonempty_rows": nonempty_rows,
            "keyword_rows": keyword_rows,
            "account_cells": account_cells[:200],
        })

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({
        "output": str(output_path),
        "sheets": [
            {
                "name": s["name"],
                "max_row": s["max_row"],
                "max_column": s["max_column"],
                "formula_count": s["formula_count"],
                "external_formula_count": len(s["external_formulas"]),
                "cross_sheet_formula_count": len(s["cross_sheet_formulas"]),
            }
            for s in result["sheets"]
        ],
        "external_links": result["external_links"],
        "account_match_count": len(result["account_matches"]),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
