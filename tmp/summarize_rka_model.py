import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


def norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip()


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def triple_key(description, lob, amount):
    if not number(amount):
        return None
    return (norm(description).casefold(), norm(lob).casefold(), round(float(amount), 2))


def main():
    rka_path = Path(sys.argv[1])
    lr_path = Path(sys.argv[2])
    output_path = Path(sys.argv[3])
    rka_f = load_workbook(rka_path, data_only=False)
    rka_v = load_workbook(rka_path, data_only=True)
    lr_v = load_workbook(lr_path, data_only=True)

    lr_by_sheet = {}
    for ws in lr_v.worksheets:
        records = []
        for r in range(1, ws.max_row + 1):
            entity = ws.cell(r, 1).value
            desc = ws.cell(r, 4).value
            lob = ws.cell(r, 6).value
            amount = ws.cell(r, 8).value
            key = triple_key(desc, lob, amount)
            if not entity or not key:
                continue
            records.append({"row": r, "desc": norm(desc), "lob": norm(lob), "amount": amount, "key": key})
        lr_by_sheet[ws.title] = records

    report = {"sheets": [], "corporate_lines": [], "formula_inconsistencies": [], "raw_reconciliation": [], "error_cells": []}

    for ws_f in rka_f.worksheets:
        ws_v = rka_v[ws_f.title]
        for row in ws_f.iter_rows():
            for cell in row:
                formula_value = cell.value
                cached_value = ws_v[cell.coordinate].value
                combined = f"{formula_value or ''} {cached_value or ''}"
                if re.search(r"#REF!|#DIV/0!|#VALUE!|#NAME\?|#N/A|#NUM!|#NULL!|#SPILL!|#CALC!", combined):
                    report["error_cells"].append({
                        "sheet": ws_f.title,
                        "cell": cell.coordinate,
                        "formula": formula_value,
                        "cached": cached_value,
                        "row_label": norm(ws_f.cell(cell.row, 5).value) or norm(ws_f.cell(cell.row, 4).value),
                    })
        lines = []
        for r in range(8, min(ws_f.max_row, 175) + 1):
            section = norm(ws_f.cell(r, 4).value)
            label = norm(ws_f.cell(r, 5).value)
            if not section and not label:
                continue
            line = {
                "row": r,
                "code": norm(ws_f.cell(r, 3).value),
                "section": section,
                "label": label,
                "budget_total": ws_v.cell(r, 11).value,
                "actual_total": ws_v.cell(r, 18).value,
                "achievement_total": ws_v.cell(r, 24).value if ws_f.title == "KORPORAT KANWIL" else ws_v.cell(r, 19).value,
                "budget_by_lob": [ws_v.cell(r, c).value for c in range(6, 11)],
                "actual_by_lob": [ws_v.cell(r, c).value for c in range(12, 18)],
                "budget_total_formula": ws_f.cell(r, 11).value if isinstance(ws_f.cell(r, 11).value, str) and ws_f.cell(r, 11).value.startswith("=") else None,
                "actual_total_formula": ws_f.cell(r, 18).value if isinstance(ws_f.cell(r, 18).value, str) and ws_f.cell(r, 18).value.startswith("=") else None,
                "achievement_formula": ws_f.cell(r, 24).value if ws_f.title == "KORPORAT KANWIL" else ws_f.cell(r, 19).value,
            }
            lines.append(line)
        if ws_f.title == "KORPORAT KANWIL":
            report["corporate_lines"] = lines

        raw_info = None
        if ws_f.title in lr_by_sheet:
            blocks = [
                ("KUR", 23, 24, 25),
                ("NON_KUR", 27, 28, 29),
                ("PEN", 31, 32, 33),
            ]
            raw_records = []
            block_stats = []
            for block_name, desc_col, lob_col, amount_col in blocks:
                block_rows = []
                for r in range(9, ws_f.max_row + 1):
                    desc = ws_v.cell(r, desc_col).value
                    lob = ws_v.cell(r, lob_col).value
                    amount = ws_v.cell(r, amount_col).value
                    key = triple_key(desc, lob, amount)
                    if not key:
                        continue
                    rec = {"block": block_name, "row": r, "desc": norm(desc), "lob": norm(lob), "amount": amount, "key": key}
                    block_rows.append(rec)
                    raw_records.append(rec)
                block_stats.append({
                    "block": block_name,
                    "rows": len(block_rows),
                    "sum": sum(rec["amount"] for rec in block_rows),
                    "distinct_descriptions": len({rec["desc"] for rec in block_rows}),
                    "distinct_lobs": len({rec["lob"] for rec in block_rows}),
                })

            lr_records = lr_by_sheet[ws_f.title]
            lr_counter = Counter(rec["key"] for rec in lr_records)
            matched = []
            unmatched_rka = []
            remaining = lr_counter.copy()
            for rec in raw_records:
                if remaining[rec["key"]] > 0:
                    remaining[rec["key"]] -= 1
                    matched.append(rec)
                else:
                    unmatched_rka.append(rec)
            unmatched_lr_count = sum(remaining.values())
            unmatched_lr = []
            remaining_copy = remaining.copy()
            for rec in lr_records:
                if remaining_copy[rec["key"]] > 0:
                    remaining_copy[rec["key"]] -= 1
                    unmatched_lr.append(rec)
            matched_sum = sum(rec["amount"] for rec in matched)
            raw_sum = sum(rec["amount"] for rec in raw_records)
            lr_sum = sum(rec["amount"] for rec in lr_records)
            raw_info = {
                "blocks": block_stats,
                "raw_rows": len(raw_records),
                "lr_rows": len(lr_records),
                "matched_rows": len(matched),
                "unmatched_rka_rows": len(unmatched_rka),
                "unmatched_lr_rows": unmatched_lr_count,
                "matched_sum": matched_sum,
                "raw_sum": raw_sum,
                "lr_sum": lr_sum,
                "sum_difference": raw_sum - lr_sum,
                "unmatched_rka_samples": [{k: v for k, v in rec.items() if k != "key"} for rec in unmatched_rka[:15]],
                "unmatched_lr_samples": [{k: v for k, v in rec.items() if k != "key"} for rec in unmatched_lr[:20]],
            }
            report["raw_reconciliation"].append({"sheet": ws_f.title, **raw_info})

        formula_rules = []
        for r in range(8, min(ws_f.max_row, 169) + 1):
            label = norm(ws_f.cell(r, 5).value) or norm(ws_f.cell(r, 4).value)
            formulas = []
            for c in range(12, 19):
                value = ws_f.cell(r, c).value
                if isinstance(value, str) and value.startswith("="):
                    formulas.append({"cell": ws_f.cell(r, c).coordinate, "formula": value})
            if formulas:
                formula_rules.append({"row": r, "label": label, "formulas": formulas})

        report["sheets"].append({
            "name": ws_f.title,
            "lines": lines,
            "raw": raw_info,
            "formula_rules": formula_rules,
        })

    branch_names = ["SURABAYA", "KEDIRI", "MALANG", "MADIUN", "BANYUWANGI"]
    formula_map = defaultdict(dict)
    for branch in branch_names:
        ws = rka_f[branch]
        for row in range(8, 170):
            label = norm(ws.cell(row, 5).value) or norm(ws.cell(row, 4).value)
            for col in range(12, 20):
                formula = ws.cell(row, col).value
                if not isinstance(formula, str) or not formula.startswith("="):
                    continue
                normalized = re.sub(r"\$?([A-Z]{1,3})\$?\d+", lambda m: f"${m.group(1)}$ROW", formula.upper())
                formula_map[(row, col, label)][branch] = normalized

    for key, by_branch in formula_map.items():
        if len(set(by_branch.values())) > 1:
            row, col, label = key
            report["formula_inconsistencies"].append({"row": row, "column": col, "label": label, "by_branch": by_branch})

    output_path.write_text(json.dumps(report, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    print(json.dumps({
        "output": str(output_path),
        "raw_reconciliation": report["raw_reconciliation"],
        "formula_inconsistency_count": len(report["formula_inconsistencies"]),
        "corporate_line_count": len(report["corporate_lines"]),
    }, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
